import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rl_canvas

OUT = Path(__file__).parent.parent / "storage" / "e1_test_fixtures"


def _pdf(path: Path, lines: list[str]) -> None:
    c = rl_canvas.Canvas(str(path), pagesize=A4)
    _, height = A4
    y = height - 60
    for line in lines:
        c.drawString(50, y, line)
        y -= 18
    c.save()
    print(f"  created {path.name}")


def _docx(path: Path, paragraphs: list[str]) -> None:
    doc = Document()
    for para in paragraphs:
        doc.add_paragraph(para)
    doc.save(path)
    print(f"  created {path.name}")


def _xlsx(path: Path, sheet_name: str, headers: list, rows: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    print(f"  created {path.name}")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    count = 0

    # 1 — SACS-002 PDF
    _pdf(OUT / "SACS-002_Third_Party_Cybersecurity_Standard.pdf", [
        "3.1 The contractor shall implement multi-factor authentication.",
        "3.2 All systems must comply with NCA ECC-2 2024.",
        "3.3 Data shall not leave the Kingdom of Saudi Arabia.",
        "3.4 Failure to comply will result in disqualification of the bid.",
    ])
    count += 1

    # 2 — BOQ xlsx
    _xlsx(
        OUT / "BOQ_4203193153.xlsx",
        "Main BOQ",
        ["Part Number", "Description", "Qty", "Unit Price (USD)", "Total Price (USD)"],
        [
            ("C9300-48U-A",   "Cisco Catalyst 9300 48-port UPOE",               5, 8500.00, 42500.00),
            ("C9300-NM-8X",   "Cisco Catalyst 9300 8x10GE Network Module",       5, 1200.00,  6000.00),
            ("DNA-A-48-3Y",   "Cisco DNA Advantage 48-port 3-year license",       5,  950.00,  4750.00),
        ],
    )
    count += 1

    # 3 — NDA docx
    _docx(OUT / "NDA_Confidentiality_Agreement.docx", [
        "Whereas the parties hereby agree to maintain confidentiality.",
        "The vendor shall not disclose any information without prior written consent from Saudi Aramco.",
        "Jurisdiction: Kingdom of Saudi Arabia. Governing law: Saudi Arabian law.",
    ])
    count += 1

    # 4 — Technical Bid Requirements PDF
    _pdf(OUT / "Technical_Bid_Requirements.pdf", [
        "4.1 Vendor shall provide 24x7 technical support with 4-hour response time.",
        "4.2 All equipment must be new and unused.",
        "4.3 Failure to submit bid bond of 2% will result in disqualification.",
        "4.4 Sealed envelope submission required by closing date 15-Jun-2026 14:00 AST.",
    ])
    count += 1

    # 5 — Scope of Work docx
    _docx(OUT / "Scope_of_Work_Network_Upgrade.docx", [
        "The scope includes campus network infrastructure upgrade across 3 sites.",
        "Vendor may propose alternative solutions where applicable.",
        "Unless otherwise specified, all documentation shall be in English and Arabic.",
        "Subject to site access approval, installation shall complete within 90 days.",
    ])
    count += 1

    # 6 — MSG (empty)
    (OUT / "Meeting_Notes_Clarifications.msg").touch()
    print("  created Meeting_Notes_Clarifications.msg")
    count += 1

    # 7 — DWG (empty)
    (OUT / "site_plan_floor1.dwg").touch()
    print("  created site_plan_floor1.dwg")
    count += 1

    # 8 — IKTVA xlsx
    _xlsx(
        OUT / "IKTVA_Local_Content_Template.xlsx",
        "IKTVA Plan",
        ["Category", "Local Content %", "Saudi Employees", "Total Employees"],
        [("Workforce", 30, 15, 50)],
    )
    count += 1

    # 9 — Commercial Registration PDF
    _pdf(OUT / "Commercial_Registration_Certificate.pdf", [
        "Authorized signatory of the company hereby certifies.",
        "Chamber of Commerce registration number: 1234567890.",
        "ZATCA VAT number: 300123456789003.",
    ])
    count += 1

    # 10 — Evaluation Criteria xlsx
    _xlsx(
        OUT / "Evaluation_Criteria_Scoring.xlsx",
        "Technical Evaluation",
        ["Criteria", "Weight %", "Min Score"],
        [
            ("System Architecture", 30, 70),
            ("Security Compliance", 25, 75),
            ("Implementation Plan", 20, 60),
            ("Experience",          15, 70),
            ("Commercial",          10,  0),
        ],
    )
    count += 1

    # manual_classification.json
    classifications = [
        {"filename": "SACS-002_Third_Party_Cybersecurity_Standard.pdf", "expected_type": "compliance",   "expected_subtype": "cybersecurity_standard", "min_confidence": 0.85},
        {"filename": "BOQ_4203193153.xlsx",                              "expected_type": "commercial",   "expected_subtype": "boq_template",          "min_confidence": 0.85},
        {"filename": "NDA_Confidentiality_Agreement.docx",               "expected_type": "legal",        "expected_subtype": "nda",                   "min_confidence": 0.85},
        {"filename": "Technical_Bid_Requirements.pdf",                   "expected_type": "technical",    "expected_subtype": "requirements",          "min_confidence": 0.75},
        {"filename": "Scope_of_Work_Network_Upgrade.docx",               "expected_type": "technical",    "expected_subtype": "requirements",          "min_confidence": 0.80},
        {"filename": "Meeting_Notes_Clarifications.msg",                 "expected_type": "correspondence","expected_subtype": "email",                 "min_confidence": 0.85},
        {"filename": "site_plan_floor1.dwg",                             "expected_type": "technical",    "expected_subtype": "engineering_drawing",   "min_confidence": 0.90},
        {"filename": "IKTVA_Local_Content_Template.xlsx",                "expected_type": "compliance",   "expected_subtype": "local_content",         "min_confidence": 0.80},
        {"filename": "Commercial_Registration_Certificate.pdf",          "expected_type": "admin",        "expected_subtype": "certificate",           "min_confidence": 0.75},
        {"filename": "Evaluation_Criteria_Scoring.xlsx",                 "expected_type": "commercial",   "expected_subtype": "evaluation_criteria",   "min_confidence": 0.80},
    ]
    json_path = OUT / "manual_classification.json"
    json_path.write_text(json.dumps(classifications, indent=2))
    print("  created manual_classification.json")

    print(f"\n{count} fixture files + 1 JSON created in {OUT}")


if __name__ == "__main__":
    main()
